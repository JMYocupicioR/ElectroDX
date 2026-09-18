"""Genera el PowerPoint 1:1 de introducción para alumnos ElectoDX."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "public" / "icons" / "icon-512x512.png"
OUT_DIR = ROOT / "docs" / "intro"
OUT_FILE = OUT_DIR / "ElectoDX-Introduccion-Alumno.pptx"

NAVY = RGBColor(0x0B, 0x13, 0x29)
NAVY_DEEP = RGBColor(0x07, 0x0C, 0x1C)
CARD = RGBColor(0x12, 0x1C, 0x38)
CARD_ALT = RGBColor(0x16, 0x24, 0x48)
COBALT = RGBColor(0x25, 0x63, 0xEB)
CYAN = RGBColor(0x06, 0xB6, 0xD4)
ACCENT = RGBColor(0x38, 0xBD, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE = RGBColor(0xCB, 0xD5, 0xE1)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
EMERALD = RGBColor(0x34, 0xD3, 0x99)
ROSE = RGBColor(0xFB, 0x71, 0x85)
LINE = RGBColor(0x1E, 0x3A, 0x5F)

FONT = "Calibri"
SLIDE_IN = 10.0
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

MODULES = [
    ("01", "Fundamentos de Neurofisiología Clínica"),
    ("02", "Estudios de Conducción Nerviosa"),
    ("03", "Electromiografía de Aguja"),
    ("04", "Respuestas Tardías y Reflejos"),
    ("05", "Estimulación Nerviosa Repetitiva"),
    ("06", "Potenciales Evocados"),
    ("07", "Estudios Especiales y Técnicas Avanzadas"),
    ("08", "Anatomía Topográfica y Neuroconducción por Nervio"),
    ("09", "Patologías Neuromusculares y Patrones Electrodiagnósticos"),
    ("10", "Criterios Diagnósticos y Algoritmos Clínicos"),
    ("11", "Referencia Rápida y Tablas Clínicas"),
    ("12", "Referencias y Bibliografía"),
    ("13", "Seguridad, Errores y Control de Calidad"),
]


def rgb(shape, color: RGBColor, line: RGBColor | None = None) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)


def set_run(run, text: str, size: int, color: RGBColor, bold: bool = False) -> None:
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    size: int,
    color: RGBColor,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    set_run(p.add_run() if p.runs else p.add_run(), text, size, color, bold)
    if p.runs:
        set_run(p.runs[0], text, size, color, bold)
    return box


def write_box(slide, left, top, width, height, lines: list[tuple[str, int, RGBColor, bool]], align=PP_ALIGN.LEFT, anchor="t"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf._txBody.bodyPr.set("anchor", anchor)
    first = True
    for text, size, color, bold in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(4)
        run = p.add_run()
        set_run(run, text, size, color, bold)
    return box


def rounded(slide, left, top, width, height, fill: RGBColor, line: RGBColor | None = None, radius=0.12):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    rgb(shape, fill, line)
    try:
        shape.adjustments[0] = radius
    except Exception:
        pass
    return shape


def oval(slide, left, top, width, height, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(width), Inches(height))
    rgb(shape, fill)
    return shape


def add_logo(slide, left, top, size) -> None:
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(left), Inches(top), Inches(size), Inches(size))


def wordmark(slide, left, top, title_size=28, show_badge=True) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(6.4), Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    r1 = p.add_run()
    set_run(r1, "Electo", title_size, WHITE, True)
    r2 = p.add_run()
    set_run(r2, "DX", title_size, ACCENT, True)
    if show_badge:
        r3 = p.add_run()
        set_run(r3, "   DIPLOMADO", 11, MUTED, True)


def footer_bar(slide, index: int, total: int) -> None:
    track = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(9.55), Inches(8.6), Inches(0.08))
    rgb(track, RGBColor(0x1A, 0x2A, 0x4A))
    fill_w = max(0.4, 8.6 * ((index + 1) / total))
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(9.55), Inches(fill_w), Inches(0.08))
    rgb(bar, CYAN)
    write_box(
        slide,
        0.7,
        9.66,
        5.5,
        0.28,
        [("ElectoDX  ·  Video de introducción", 10, MUTED, False)],
    )
    write_box(
        slide,
        6.3,
        9.66,
        3.0,
        0.28,
        [(f"{index + 1:02d}  /  {total:02d}", 10, MUTED, True)],
        align=PP_ALIGN.RIGHT,
    )


def header(slide, kicker: str) -> None:
    add_logo(slide, 0.55, 0.38, 0.52)
    wordmark(slide, 1.18, 0.42, 18, show_badge=True)
    write_box(slide, 0.55, 1.05, 8.9, 0.28, [(kicker.upper(), 11, CYAN, True)])


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text.strip()


def fade_transition(slide) -> None:
    sld = slide._element
    for child in list(sld):
        if child.tag == qn("p:transition") or child.tag.endswith("transition"):
            sld.remove(child)
    tr = etree.SubElement(sld, f"{{{P_NS}}}transition")
    tr.set("spd", "med")
    etree.SubElement(tr, f"{{{P_NS}}}fade")


def add_emg(slide, left, top, width, height) -> None:
    pts = [
        (0.00, 0.55),
        (0.10, 0.55),
        (0.14, 0.18),
        (0.18, 0.82),
        (0.22, 0.55),
        (0.38, 0.55),
        (0.46, 0.20),
        (0.54, 0.12),
        (0.64, 0.80),
        (0.72, 0.55),
        (1.00, 0.55),
    ]
    x0 = Inches(left + pts[0][0] * width)
    y0 = Inches(top + pts[0][1] * height)
    builder = slide.shapes.build_freeform(Emu(x0), Emu(y0))
    segs = [
        (Emu(Inches(left + x * width)), Emu(Inches(top + y * height)))
        for x, y in pts[1:]
    ]
    builder.add_line_segments(segs, close=False)
    shape = builder.convert_to_shape()
    shape.fill.background()
    shape.line.color.rgb = ACCENT
    shape.line.width = Pt(2.4)


def chrome(slide, left, top, width, height, url: str):
    rounded(slide, left, top, width, height, CARD, LINE, 0.08)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.38)
    )
    rgb(bar, RGBColor(0x0D, 0x16, 0x2C), LINE)
    for i, color in enumerate((ROSE, AMBER, EMERALD)):
        oval(slide, left + 0.12 + i * 0.18, top + 0.12, 0.14, 0.14, color)
    write_box(slide, left + 0.72, top + 0.06, width - 0.9, 0.28, [(url, 9, MUTED, False)])
    return top + 0.46


def new_slide(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(SLIDE_IN), Inches(SLIDE_IN))
    rgb(bg, NAVY)
    # ambient orbs
    o1 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1.2), Inches(-1.4), Inches(4.2), Inches(4.2))
    rgb(o1, RGBColor(0x14, 0x2A, 0x5C))
    o2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.2), Inches(7.4), Inches(4.0), Inches(4.0))
    rgb(o2, RGBColor(0x0C, 0x3A, 0x48))
    return slide


def card_stat(slide, left, top, w, h, value: str, label: str) -> None:
    rounded(slide, left, top, w, h, CARD, LINE, 0.12)
    write_box(slide, left + 0.18, top + 0.22, w - 0.3, 0.55, [(value, 28, WHITE, True)])
    write_box(slide, left + 0.18, top + 0.78, w - 0.3, 0.4, [(label, 12, MUTED, False)])


def pill(slide, left, top, w, h, text: str, fill=COBALT, fg=WHITE) -> None:
    rounded(slide, left, top, w, h, fill, None, 0.5)
    write_box(slide, left, top + 0.04, w, h - 0.04, [(text, 11, fg, True)], align=PP_ALIGN.CENTER, anchor="ctr")


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_IN)
    prs.slide_height = Inches(SLIDE_IN)
    slides_meta: list[tuple] = []

    # 1. Portada
    s = new_slide(prs)
    add_logo(s, 3.85, 1.55, 2.3)
    wordmark(s, 1.7, 4.05, 40, True)
    write_box(
        s,
        1.2,
        4.85,
        7.6,
        0.9,
        [
            ("Plataforma integral de formación en", 16, SLATE, False),
            ("electrodiagnóstico y neurofisiología clínica", 16, SLATE, False),
        ],
        align=PP_ALIGN.CENTER,
    )
    pill(s, 2.7, 6.05, 4.6, 0.42, "VIDEO DE INTRODUCCIÓN AL ALUMNO", RGBColor(0x0E, 0x3A, 0x4A), ACCENT)
    add_emg(s, 1.3, 6.7, 7.4, 1.35)
    write_box(s, 1.5, 8.2, 7.0, 0.4, [("Lee las notas del presentador · un clic = una diapositiva", 12, MUTED, False)], align=PP_ALIGN.CENTER)
    notes(
        s,
        "Bienvenido, bienvenida. Si estás viendo este video es porque acabas de entrar a ElectoDX, "
        "la plataforma de formación en electrodiagnóstico y neurofisiología clínica. "
        "Mi objetivo en los próximos minutos es simple: que cuando termines sepas exactamente qué vas a aprender, "
        "cómo se estudia aquí, cómo se te va a evaluar y qué necesitas cumplir para obtener tu constancia. "
        "Te recomiendo verlo completo una sola vez, con calma.",
    )
    slides_meta.append(s)

    # 2. Agenda
    s = new_slide(prs)
    header(s, "Antes de empezar")
    write_box(s, 0.55, 1.35, 8.9, 1.1, [("Cuatro cosas que debes tener claras", 28, WHITE, True)])
    items = [
        ("01", "En qué consiste el diplomado"),
        ("02", "Cómo entrar y navegar la plataforma"),
        ("03", "Cómo se estudia y cómo se evalúa"),
        ("04", "Cómo se obtiene y verifica la constancia"),
    ]
    for i, (num, label) in enumerate(items):
        top = 2.65 + i * 1.35
        rounded(s, 0.55, top, 8.9, 1.18, CARD, LINE, 0.1)
        oval(s, 0.82, top + 0.28, 0.62, 0.62, COBALT)
        write_box(s, 0.82, top + 0.38, 0.62, 0.45, [(num, 14, WHITE, True)], align=PP_ALIGN.CENTER)
        write_box(s, 1.7, top + 0.35, 7.3, 0.5, [(label, 20, WHITE, True)])
    notes(
        s,
        "Nada de sorpresas a mitad del camino. En este video cubrimos el programa, tu acceso, "
        "la navegación diaria, las evaluaciones y el cierre con la constancia.",
    )
    slides_meta.append(s)

    # 3. Qué es el curso
    s = new_slide(prs)
    header(s, "El curso")
    write_box(s, 0.55, 1.35, 8.9, 1.3, [("Un programa completo,\nno clases sueltas", 30, WHITE, True)])
    write_box(
        s,
        0.55,
        2.8,
        8.9,
        0.7,
        [("Temario jerárquico con perlas clínicas, puntos clave, bibliografía y práctica.", 15, SLATE, False)],
    )
    stats = [("13", "Módulos"), ("479", "Temas y subtemas"), ("77 mil", "Palabras"), ("378", "Lecciones redactadas")]
    for i, (n, l) in enumerate(stats):
        card_stat(s, 0.55 + (i % 2) * 4.5, 3.7 + (i // 2) * 2.35, 4.25, 2.1, n, l)
    notes(
        s,
        "Esto no es un curso de video-clases sueltas. Es un programa de posgrado estructurado en 13 módulos, "
        "con un temario de casi quinientos temas y subtemas, y alrededor de setenta y siete mil palabras de contenido "
        "redactado con bibliografía.",
    )
    slides_meta.append(s)

    # 4. Módulos
    s = new_slide(prs)
    header(s, "Temario oficial")
    write_box(s, 0.55, 1.32, 8.9, 0.5, [("Los 13 módulos", 28, WHITE, True)])
    for i, (num, title) in enumerate(MODULES):
        col = 0 if i < 7 else 1
        row = i if i < 7 else i - 7
        left = 0.5 + col * 4.75
        top = 1.95 + row * 0.95
        rounded(s, left, top, 4.55, 0.85, CARD, LINE, 0.1)
        write_box(s, left + 0.12, top + 0.12, 0.55, 0.28, [(num, 11, CYAN, True)])
        write_box(s, left + 0.12, top + 0.38, 4.25, 0.4, [(title, 11, WHITE, True)])
    notes(
        s,
        "El recorrido está diseñado en el mismo orden en el que se construye un electromiografista: "
        "bases, técnicas, anatomía topográfica, patología, algoritmos, tablas, bibliografía y un módulo "
        "completo de seguridad, errores frecuentes y control de calidad. Ese último no lo dejes para el final.",
    )
    slides_meta.append(s)

    # 5. Recorrido
    s = new_slide(prs)
    header(s, "Ruta de aprendizaje")
    write_box(s, 0.55, 1.35, 8.9, 1.0, [("El mismo orden de un electromiografista", 26, WHITE, True)])
    steps = [
        ("01  Bases", "Electricidad, neuroanatomía funcional y fisiología de la contracción muscular."),
        ("02  Técnicas", "NCS, EMG de aguja, respuestas tardías, estimulación repetitiva y potenciales evocados."),
        ("03  Anatomía", "Neuroconducción nervio por nervio. Aquí se resuelve la mayoría de los estudios reales."),
        ("04  Clínica", "Patrones, algoritmos, tablas y control de calidad. El día que estés solo frente al paciente."),
    ]
    for i, (t, d) in enumerate(steps):
        top = 2.55 + i * 1.5
        rounded(s, 0.55, top, 8.9, 1.32, CARD, LINE, 0.1)
        stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(top), Inches(0.12), Inches(1.32))
        rgb(stripe, CYAN if i % 2 == 0 else COBALT)
        write_box(s, 0.95, top + 0.18, 8.2, 0.4, [(t, 18, WHITE, True)])
        write_box(s, 0.95, top + 0.62, 8.2, 0.5, [(d, 13, SLATE, False)])
    notes(
        s,
        "Primero las bases: sin esto, todo lo demás son recetas memorizadas. Después las técnicas. "
        "Luego la anatomía topográfica nervio por nervio. Y finalmente patología, algoritmos y control de calidad.",
    )
    slides_meta.append(s)

    # 6. Acceso
    s = new_slide(prs)
    header(s, "Tu acceso")
    write_box(s, 0.55, 1.32, 8.9, 0.7, [("Primero tu expediente médico", 28, WHITE, True)])
    inner = chrome(s, 0.5, 2.15, 4.55, 5.55, "electodx.app/auth/registro")
    write_box(s, 0.7, inner + 0.1, 4.15, 0.3, [("Paso 1 de 3  ·  Datos de cuenta", 10, CYAN, True)])
    write_box(s, 0.7, inner + 0.45, 4.15, 0.3, [("Validación oficial SEP México", 12, WHITE, True)])
    rounded(s, 0.7, inner + 0.9, 4.15, 1.55, CARD_ALT, LINE, 0.1)
    write_box(s, 0.88, inner + 1.02, 3.8, 0.25, [("Cédula profesional", 10, MUTED, False)])
    rounded(s, 0.88, inner + 1.32, 3.8, 0.38, RGBColor(0x0B, 0x13, 0x29), LINE, 0.08)
    write_box(s, 1.0, inner + 1.36, 3.5, 0.3, [("12345678", 12, SLATE, False)])
    pill(s, 0.88, inner + 1.82, 3.8, 0.4, "Verificar cédula", COBALT, WHITE)
    rounded(s, 0.7, inner + 2.65, 4.15, 0.7, RGBColor(0x0C, 0x3A, 0x2E), RGBColor(0x34, 0xD3, 0x99), 0.1)
    write_box(s, 0.88, inner + 2.82, 3.8, 0.4, [("Cédula verificada oficialmente", 12, EMERALD, True)])
    facts = [
        ("Cédula SEP", "Opcional al registrarte si eres R1 o está en trámite. Obligatoria para emitir la constancia."),
        ("Perfil formativo", "Residente de rehabilitación, especialista, neurofisiólogo o médico de especialidad afín."),
        ("Admisión", "Tras el registro, tu expediente pasa a revisión. Mientras dice “En espera de admisión” ves el temario, no los módulos."),
    ]
    for i, (t, d) in enumerate(facts):
        top = 2.15 + i * 1.85
        rounded(s, 5.25, top, 4.25, 1.7, CARD, LINE, 0.1)
        write_box(s, 5.45, top + 0.18, 3.9, 0.35, [(t, 14, CYAN, True)])
        write_box(s, 5.45, top + 0.55, 3.9, 0.95, [(d, 12, SLATE, False)])
    notes(
        s,
        "Al registrarte te pedimos nombre, correo, perfil formativo, sede hospitalaria y universidad. "
        "La cédula se valida contra el Registro Nacional de Profesionistas de la SEP. "
        "Si estás en trámite puedes registrarte sin ella, pero resuélvela durante el curso: "
        "la cédula verificada es requisito para la constancia. "
        "Después, tu expediente queda en revisión académica.",
    )
    slides_meta.append(s)

    # 7. Navegación
    s = new_slide(prs)
    header(s, "Cómo te mueves")
    write_box(s, 0.55, 1.32, 8.9, 0.6, [("Cinco puertas, un mismo techo", 28, WHITE, True)])
    nav = [
        ("Temario", "Índice completo del programa. Tu mapa general."),
        ("Currículo", "El árbol de 13 módulos. Úsalo todos los días. Tiene buscador."),
        ("Exámenes", "Simulacros tipo consejo, con o sin cronómetro."),
        ("Simuladores", "Práctica interactiva: plexo, casos, trazos."),
        ("Mi Portal", "Tu casa: progreso, tareas, constancia y estudio."),
    ]
    inner = chrome(s, 0.5, 2.1, 9.0, 1.55, "electodx.app")
    labels = ["Temario", "Simuladores  PRO", "Exámenes", "Currículo", "Mi Portal"]
    for i, lab in enumerate(labels):
        fill = COBALT if i == 0 else CARD_ALT
        pill(s, 0.7 + i * 1.72, inner + 0.28, 1.6, 0.38, lab, fill, WHITE)
    for i, (t, d) in enumerate(nav):
        col, row = i % 2, i // 2
        if i == 4:
            left, top, w = 0.5, 6.95, 9.0
        else:
            left, top, w = 0.5 + col * 4.55, 3.9 + row * 1.45, 4.35
        rounded(s, left, top, w, 1.3, CARD, LINE, 0.1)
        write_box(s, left + 0.22, top + 0.2, w - 0.4, 0.35, [(t, 16, WHITE, True)])
        write_box(s, left + 0.22, top + 0.6, w - 0.4, 0.5, [(d, 13, SLATE, False)])
    notes(
        s,
        "Arriba, en la barra, tienes Temario, Currículo, Exámenes, Simuladores y Mi Portal. "
        "Currículo abre el panel lateral con el árbol completo, tu progreso y un buscador. "
        "Si necesitas encontrar onda F o túnel carpiano, escríbelo ahí. "
        "También hay selector de idioma, modo claro y oscuro, e indicador de conexión.",
    )
    slides_meta.append(s)

    # 8. Portal
    s = new_slide(prs)
    header(s, "Mi Portal")
    write_box(s, 0.55, 1.32, 8.9, 0.55, [("Empieza siempre aquí", 28, WHITE, True)])
    inner = chrome(s, 0.5, 2.0, 9.0, 6.95, "electodx.app/portal")
    rounded(s, 0.7, inner + 0.15, 8.6, 1.35, RGBColor(0x14, 0x3A, 0x6A), LINE, 0.1)
    write_box(s, 0.95, inner + 0.28, 8.2, 0.25, [("PORTAL ACADÉMICO DEL ALUMNO", 10, CYAN, True)])
    write_box(s, 0.95, inner + 0.55, 8.2, 0.4, [("Bienvenido, Dr(a).", 22, WHITE, True)])
    write_box(s, 0.95, inner + 1.0, 8.2, 0.3, [("Racha de estudio  ·  Cédula verificada SEP", 12, SLATE, False)])
    tabs = ["Resumen", "Módulos", "Quizzes", "Tareas", "Avisos", "Constancia", "Estudio"]
    for i, tab in enumerate(tabs):
        fill = COBALT if i == 0 else CARD_ALT
        pill(s, 0.7 + i * 1.22, inner + 1.7, 1.15, 0.36, tab, fill, WHITE)
    kpis = [("Progreso", "— %"), ("Promedio", "— %"), ("Horas", "CME")]
    for i, (l, v) in enumerate(kpis):
        rounded(s, 0.7 + i * 2.95, inner + 2.25, 2.8, 1.35, CARD_ALT, LINE, 0.1)
        write_box(s, 0.9 + i * 2.95, inner + 2.4, 2.45, 0.55, [(v, 26, ACCENT, True)], align=PP_ALIGN.CENTER)
        write_box(s, 0.9 + i * 2.95, inner + 3.05, 2.45, 0.35, [(l, 12, MUTED, False)], align=PP_ALIGN.CENTER)
    pill(s, 0.7, inner + 3.85, 8.6, 0.5, "Continuar donde te quedaste", COBALT, WHITE)
    write_box(
        s,
        0.7,
        inner + 4.55,
        8.6,
        0.7,
        [("La plataforma recuerda el último tema. No pierdas tiempo buscando en qué ibas.", 14, SLATE, False)],
        align=PP_ALIGN.CENTER,
    )
    notes(
        s,
        "Entra a Mi Portal. Verás progreso global, promedio, racha y el botón Continuar donde te quedaste. "
        "Las pestañas son Resumen General, Mis Clases y Módulos, Quizzes del Curso, Tareas, Avisos, "
        "Constancia, y Estudio y reportes. Esa última es la que más puede subir tu calificación.",
    )
    slides_meta.append(s)

    # 9. Estudio
    s = new_slide(prs)
    header(s, "Cómo estudiar")
    write_box(s, 0.55, 1.32, 8.9, 0.7, [("No salgas de la lección sin tus tarjetas", 24, WHITE, True)])
    tools = [
        ("Apuntes", "Quedan ligados al tema, en tu cuenta. Los retomas en cualquier dispositivo."),
        ("Marcador", "Guarda la lección como favorita para el repaso antes de un examen."),
        ("Tarjetas SM-2", "Se crean desde perlas clínicas y puntos clave. Las calificas: Otra vez, Difícil, Bien o Fácil."),
        ("Cuaderno de errores", "Junta las preguntas que fallaste y te enlaza de regreso al tema. Si solo haces una cosa antes del examen, haz esto."),
    ]
    for i, (t, d) in enumerate(tools):
        col, row = i % 2, i // 2
        left, top = 0.5 + col * 4.55, 2.2 + row * 3.2
        rounded(s, left, top, 4.35, 2.95, CARD, LINE, 0.1)
        oval(s, left + 0.25, top + 0.3, 0.55, 0.55, COBALT if i % 2 == 0 else RGBColor(0x0E, 0x5C, 0x66))
        write_box(s, left + 0.28, top + 0.4, 0.5, 0.35, [(f"{i+1:02d}", 12, WHITE, True)], align=PP_ALIGN.CENTER)
        write_box(s, left + 0.25, top + 1.05, 3.9, 0.45, [(t, 18, WHITE, True)])
        write_box(s, left + 0.25, top + 1.55, 3.9, 1.15, [(d, 13, SLATE, False)])
    notes(
        s,
        "Dentro de cada lección tienes apuntes, marcador y Crear tarjetas de este tema. "
        "En Estudio y reportes está la repetición espaciada SM-2, fortalezas por módulo, "
        "cuaderno de errores, reportes EMG, preguntas al docente y calendario exportable.",
    )
    slides_meta.append(s)

    # 10. Evaluaciones
    s = new_slide(prs)
    header(s, "Evaluaciones por tema")
    write_box(s, 0.55, 1.32, 8.9, 0.55, [("Aprueba con 70. Apunta al 80.", 28, WHITE, True)])
    inner = chrome(s, 0.5, 2.0, 5.15, 6.95, "electodx.app  ·  Evaluación del tema")
    write_box(s, 0.7, inner + 0.1, 4.75, 0.25, [("Pregunta  3  de  4        02:41", 11, MUTED, False)])
    track = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(inner + 0.45), Inches(4.75), Inches(0.08))
    rgb(track, RGBColor(0x1A, 0x2A, 0x4A))
    bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(inner + 0.45), Inches(3.4), Inches(0.08))
    rgb(bar, CYAN)
    write_box(
        s,
        0.7,
        inner + 0.7,
        4.75,
        1.15,
        [("En un CMAP del mediano, una latencia distal de 5.9 ms sugiere principalmente:", 14, WHITE, True)],
    )
    opts = [
        (False, "Variante anatómica de Martin-Gruber"),
        (True, "Desmielinización focal en muñeca"),
        (False, "Miopatía inflamatoria"),
        (False, "Lesión de neurona motora superior"),
    ]
    for i, (on, t) in enumerate(opts):
        top = inner + 2.05 + i * 0.7
        rounded(s, 0.7, top, 4.75, 0.6, RGBColor(0x0C, 0x3A, 0x4A) if on else CARD_ALT, CYAN if on else LINE, 0.1)
        write_box(s, 0.9, top + 0.14, 4.4, 0.35, [(t, 12, WHITE if on else SLATE, True if on else False)])
    facts = [
        ("70%", "Umbral para aprobar el tema. Configurable, este es el valor por defecto."),
        ("Aleatorio", "Preguntas y opciones se mezclan. No memorices posiciones."),
        ("Sin corte", "El cronómetro solo registra duración. Tómate tu tiempo."),
        ("Al terminar", "Ves la respuesta correcta y la perla clínica oficial. Léelas todas."),
        ("Reintentos", "Puedes volver a intentar. Al aprobar, el tema se marca completado."),
    ]
    for i, (t, d) in enumerate(facts):
        top = 2.0 + i * 1.38
        rounded(s, 5.85, top, 3.65, 1.25, CARD, LINE, 0.1)
        write_box(s, 6.05, top + 0.12, 3.3, 0.32, [(t, 16, CYAN, True)])
        write_box(s, 6.05, top + 0.48, 3.3, 0.65, [(d, 11, SLATE, False)])
    notes(
        s,
        "Al final de los temas evaluables está la evaluación. Se aprueba con 70 por ciento o más. "
        "Hay opción única, múltiple, verdadero o falso e identificación de imagen. "
        "Debes responder todas las preguntas. Las respuestas se califican en el servidor. "
        "Aprobar con 70 desbloquea el tema; el promedio de 80 es el que mira la constancia. "
        "No uses los reintentos para adivinar.",
    )
    slides_meta.append(s)

    # 11. Exámenes
    s = new_slide(prs)
    header(s, "Simulacros y tareas")
    write_box(s, 0.55, 1.32, 8.9, 0.7, [("Practica en Tutor. Mídete en Examen.", 26, WHITE, True)])
    blocks = [
        ("Tú eliges", "10, 20, 50 o todas las preguntas. Cronómetro opcional de 15 a 90 minutos."),
        ("Modo Tutor", "Te dice si acertaste inmediatamente. Úsalo para aprender."),
        ("Modo Examen", "No ves nada hasta el final. Condiciones parecidas a las del consejo."),
        ("No acreditan tema", "Los simulacros del hub no marcan lecciones del temario como completadas."),
    ]
    for i, (t, d) in enumerate(blocks):
        col, row = i % 2, i // 2
        left, top = 0.5 + col * 4.55, 2.2 + row * 2.35
        rounded(s, left, top, 4.35, 2.15, CARD, LINE, 0.1)
        write_box(s, left + 0.25, top + 0.28, 3.9, 0.45, [(t, 18, WHITE, True)])
        write_box(s, left + 0.25, top + 0.85, 3.9, 1.0, [(d, 14, SLATE, False)])
    rounded(s, 0.5, 7.05, 9.0, 1.95, RGBColor(0x3A, 0x2A, 0x0C), RGBColor(0xFB, 0xBF, 0x24), 0.1)
    write_box(s, 0.8, 7.25, 8.4, 0.4, [("Examen asignado por el profesor", 16, AMBER, True)])
    write_box(
        s,
        0.8,
        7.7,
        8.4,
        1.05,
        [
            (
                "Aparece en Tareas. El cronómetro sigue aunque cierres la pestaña. Al llegar a cero se envía solo. No lo abras “para echar un ojo”.",
                14,
                SLATE,
                False,
            )
        ],
    )
    notes(
        s,
        "Aparte de las evaluaciones de tema, en Exámenes armas tu propio simulacro. "
        "Modo Tutor para aprender, Modo Examen para medirte. "
        "Si tu profesor te asigna un examen, es evaluación oficial con tiempo continuo.",
    )
    slides_meta.append(s)

    # 12. Simuladores
    s = new_slide(prs)
    header(s, "Práctica")
    write_box(s, 0.55, 1.32, 8.9, 0.55, [("Leer trazos, no solo memorizarlos", 26, WHITE, True)])
    sims = [
        ("Mapa del plexo braquial", "Algoritmo de 5 pasos: raíces C5–T1, troncos y cordones."),
        ("Casos clínicos EMG", "Más de 30 casos: historia, NCS, aguja, diagnóstico y retroalimentación."),
        ("Simulador de trazos", "Compara CMAP y SNAP normal vs patológico, lado a lado."),
        ("Examen tipo consejo", "Simulacro cronometrado con análisis de brechas."),
    ]
    for i, (t, d) in enumerate(sims):
        col, row = i % 2, i // 2
        left, top = 0.5 + col * 4.55, 2.05 + row * 2.55
        rounded(s, left, top, 4.35, 2.35, CARD, LINE, 0.1)
        pill(s, left + 0.22, top + 0.22, 0.7, 0.32, "PRO", RGBColor(0x3A, 0x2A, 0x0C), AMBER)
        write_box(s, left + 0.22, top + 0.7, 3.95, 0.55, [(t, 16, WHITE, True)])
        write_box(s, left + 0.22, top + 1.3, 3.95, 0.8, [(d, 13, SLATE, False)])
    write_box(
        s,
        0.55,
        7.3,
        8.9,
        1.6,
        [
            ("Incluye osciloscopio de NCS, audio de EMG de aguja, mapa de miotomas", 14, SLATE, False),
            ("y casos con trampas técnicas: los que más enseñan.", 14, SLATE, False),
        ],
        align=PP_ALIGN.CENTER,
    )
    notes(
        s,
        "En Simuladores tienes el mapa del plexo, casos clínicos EMG, simulador de trazos y examen tipo consejo. "
        "Los casos cubren patrones axonales, desmielinizantes, miopáticos, atrapamientos, radiculopatías, "
        "plexopatías, motoneurona, unión neuromuscular y trampas técnicas.",
    )
    slides_meta.append(s)

    # 13. Constancia
    s = new_slide(prs)
    header(s, "Cierre académico")
    write_box(s, 0.55, 1.32, 8.9, 0.55, [("Cuatro luces en verde", 30, WHITE, True)])
    reqs = [
        ("01", "Cédula verificada", "Validada contra el Registro Nacional de Profesionistas de la SEP."),
        ("02", "Avance curricular", "Debes completar prácticamente todo el temario."),
        ("03", "Evaluaciones", "La mayoría de las evaluaciones disponibles, aprobadas."),
        ("04", "Promedio ≥ 80%", "El 70 aprueba el tema. El 80 es el listón de la constancia."),
    ]
    for i, (n, t, d) in enumerate(reqs):
        col, row = i % 2, i // 2
        left, top = 0.5 + col * 4.55, 2.1 + row * 2.7
        rounded(s, left, top, 4.35, 2.5, CARD, LINE, 0.1)
        oval(s, left + 0.25, top + 0.3, 0.55, 0.55, EMERALD)
        write_box(s, left + 0.25, top + 0.4, 0.55, 0.38, [(n, 11, NAVY, True)], align=PP_ALIGN.CENTER)
        write_box(s, left + 0.25, top + 1.05, 3.9, 0.45, [(t, 16, WHITE, True)])
        write_box(s, left + 0.25, top + 1.55, 3.9, 0.7, [(d, 13, SLATE, False)])
    notes(
        s,
        "Ve a Mi Portal, pestaña Constancia. Los cuatro indicadores tienen que estar en verde: "
        "cédula verificada, avance curricular prácticamente al total, mayoría de evaluaciones aprobadas "
        "y promedio de al menos 80 por ciento. Apunta más alto desde el principio. "
        "Cuando estén cumplidos se activa Emitir constancia.",
    )
    slides_meta.append(s)

    # 14. Verificación
    s = new_slide(prs)
    header(s, "Documento verificable")
    write_box(s, 0.55, 1.32, 8.9, 0.9, [("No es un PDF\nque se pueda inventar", 28, WHITE, True)])
    inner = chrome(s, 2.15, 2.4, 5.7, 5.0, "electodx.app/verificar/EDX-2026-A9F3C1")
    add_logo(s, 4.35, inner + 0.25, 1.3)
    oval(s, 4.55, inner + 1.7, 0.9, 0.9, RGBColor(0x0C, 0x3A, 0x2E))
    write_box(s, 2.4, inner + 2.7, 5.2, 0.45, [("Constancia vigente", 20, EMERALD, True)], align=PP_ALIGN.CENTER)
    write_box(s, 2.4, inner + 3.2, 5.2, 0.35, [("Folio EDX-AÑO-código  ·  QR en el PDF", 13, SLATE, False)], align=PP_ALIGN.CENTER)
    write_box(
        s,
        0.7,
        7.7,
        8.6,
        1.2,
        [
            ("Se emite una sola vez. Si la pides de nuevo, es el mismo folio.", 15, SLATE, False),
            ("Cualquier hospital puede confirmarla en segundos con el folio o el QR.", 15, SLATE, False),
        ],
        align=PP_ALIGN.CENTER,
    )
    notes(
        s,
        "La plataforma genera un folio único, tu promedio, un código de verificación y un código QR. "
        "Cualquier jefe de servicio puede entrar a la dirección de verificación con tu folio o escanear el QR. "
        "Si vuelves a pedirla, te entregamos la misma, con el mismo folio.",
    )
    slides_meta.append(s)

    # 15. Tips
    s = new_slide(prs)
    header(s, "Detalles que importan")
    write_box(s, 0.55, 1.32, 8.9, 0.55, [("Cinco hábitos de hospital", 28, WHITE, True)])
    tips = [
        ("01", "Instálala como app", "Agrégala a la pantalla de inicio. Se abre a pantalla completa, como app nativa."),
        ("02", "Activa notificaciones", "Así te enteras cuando asignen un examen, un caso o una calificación."),
        ("03", "Descarga módulos", "Puedes consultarlos sin internet si el servicio tiene mala señal."),
        ("04", "Tu progreso es de la cuenta", "Empiezas en la computadora del servicio y sigues en el teléfono."),
        ("05", "La asistencia cuenta", "No entra en la constancia. Sí forma parte del kardex: clases y talleres en vivo."),
    ]
    for i, (n, t, d) in enumerate(tips):
        top = 2.05 + i * 1.35
        rounded(s, 0.5, top, 9.0, 1.22, CARD, LINE, 0.1)
        oval(s, 0.75, top + 0.3, 0.6, 0.6, COBALT)
        write_box(s, 0.75, top + 0.42, 0.6, 0.38, [(n[-2:], 12, WHITE, True)], align=PP_ALIGN.CENTER)
        write_box(s, 1.6, top + 0.18, 7.6, 0.35, [(t, 16, WHITE, True)])
        write_box(s, 1.6, top + 0.58, 7.6, 0.45, [(d, 13, SLATE, False)])
    notes(
        s,
        "Instala la plataforma como aplicación. Activa las notificaciones. Descarga módulos para estudiar sin red. "
        "Todo el progreso vive en tu cuenta, no en el dispositivo. "
        "La asistencia a clases y talleres cuenta en el expediente académico, no para emitir la constancia.",
    )
    slides_meta.append(s)

    # 16. Ruta corta
    s = new_slide(prs)
    header(s, "Si empezarás hoy")
    write_box(s, 0.55, 1.32, 8.9, 0.7, [("Un ciclo. El avance llega solo.", 28, WHITE, True)])
    path = [
        "Abre Mi Portal",
        "Continúa el tema",
        "Lee perlas y puntos clave",
        "Genera tus tarjetas",
        "Resuelve la evaluación",
        "Lee todas las explicaciones",
        "Cada semana: cuaderno de errores",
        "Cada semana: tarjetas vencidas",
    ]
    for i, t in enumerate(path):
        col, row = i % 2, i // 2
        left, top = 0.5 + col * 4.55, 2.2 + row * 1.6
        rounded(s, left, top, 4.35, 1.4, CARD, LINE, 0.1)
        write_box(s, left + 0.25, top + 0.22, 3.9, 0.3, [(f"{i+1:02d}", 12, CYAN, True)])
        write_box(s, left + 0.25, top + 0.6, 3.9, 0.55, [(t, 16, WHITE, True)])
    notes(
        s,
        "La ruta corta: entra a tu portal, usa Continuar donde te quedaste, estudia el tema completo, "
        "genera tarjetas, contesta la evaluación y lee todas las explicaciones, aciertos incluidos. "
        "Cada semana, cuaderno de errores y repasos pendientes.",
    )
    slides_meta.append(s)

    # 17. Cierre
    s = new_slide(prs)
    add_logo(s, 3.95, 1.7, 2.1)
    wordmark(s, 1.7, 4.05, 36, True)
    write_box(
        s,
        1.1,
        4.95,
        7.8,
        1.5,
        [
            ("El electrodiagnóstico bien hecho", 18, SLATE, False),
            ("evita cirugías innecesarias.", 22, WHITE, True),
            ("El mal hecho las provoca.", 18, SLATE, False),
        ],
        align=PP_ALIGN.CENTER,
    )
    pill(s, 2.55, 6.7, 4.9, 0.48, "NOS VEMOS EN EL MÓDULO 01", COBALT, WHITE)
    write_box(
        s,
        1.2,
        7.5,
        7.6,
        0.9,
        [
            ("Si tienes dudas, usa Preguntas al docente", 14, MUTED, False),
            ("dentro de Mi Portal. Para eso estamos.", 14, MUTED, False),
        ],
        align=PP_ALIGN.CENTER,
    )
    notes(
        s,
        "Aquí no estamos memorizando datos para pasar un examen. Estamos aprendiendo a leer señales eléctricas "
        "de las que dependen decisiones quirúrgicas, tratamientos inmunológicos y diagnósticos que cambian la vida "
        "de un paciente. Estudia con esa responsabilidad. Bienvenido a ElectoDX. Nos vemos en el módulo uno.",
    )
    slides_meta.append(s)

    total = len(slides_meta)
    for i, slide in enumerate(slides_meta):
        footer_bar(slide, i, total)
        fade_transition(slide)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_FILE)
    return OUT_FILE


if __name__ == "__main__":
    path = build()
    print(path)
